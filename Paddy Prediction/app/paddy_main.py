import os
import glob
import shutil
from typing import Optional
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import seaborn as sns
import pandas as pd
import numpy as np

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from sklearn.model_selection import train_test_split, KFold, cross_val_score
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.linear_model import Ridge

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import ScatterChart, BarChart, Reference, Series

app = FastAPI(title="Paddy ML Prediction API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "Data")
OUTPUT_DIR = os.path.join(BASE_DIR, "Output")
PLOTS_DIR = os.path.join(BASE_DIR, "Plots")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(PLOTS_DIR, exist_ok=True)

app.mount("/plots", StaticFiles(directory=PLOTS_DIR), name="plots")


class ExcelWriterML:
    def __init__(self, data_file='paddydataset.csv', target=None, test_size=0.2, random_state=42):
        self.file_path = os.path.join(DATA_DIR, data_file)
        self.output_excel = os.path.join(OUTPUT_DIR, 'ml_model_report.xlsx')
        self.plots_dir = PLOTS_DIR
        self.target = target
        self.test_size = test_size
        self.random_state = random_state

        self.df = None
        self.X = None
        self.y = None
        self.X_train = None
        self.X_test = None
        self.y_train = None
        self.y_test = None
        self.pipeline = None
        self.train_pred = None
        self.test_pred = None
        self.cv_r2 = None
        self.summary_df = None
        self.eda_missing_df = None
        self.eda_stats_df = None
        self.actual_pred_df = None
        self.output_metrics_df = None
        self.feature_importance_df = None
        self.numeric_df = None
        self.plot_files = []

    def load_data(self):
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"File not found: {self.file_path}")
        self.df = pd.read_csv(self.file_path)
        self.df.columns = [str(c).strip() for c in self.df.columns]

        if not self.target or self.target not in self.df.columns:
            raise ValueError(f"Target '{self.target}' not found in columns: {list(self.df.columns)}")

        self.df[self.target] = pd.to_numeric(self.df[self.target], errors='coerce')
        self.df = self.df.dropna(subset=[self.target]).reset_index(drop=True)

        self.X = self.df.drop(columns=[self.target])
        self.y = self.df[self.target].astype(float)
        self.numeric_df = self.df.select_dtypes(include=np.number)

    def build_model(self):
        num_cols = [c for c in self.X.columns if pd.api.types.is_numeric_dtype(self.X[c])]
        cat_cols = [c for c in self.X.columns if c not in num_cols]

        transformers = []
        if num_cols:
            transformers.append(('num', Pipeline([
                ('imputer', SimpleImputer(strategy='median')),
                ('scaler', StandardScaler())
            ]), num_cols))
        if cat_cols:
            transformers.append(('cat', Pipeline([
                ('imputer', SimpleImputer(strategy='most_frequent')),
                ('onehot', OneHotEncoder(handle_unknown='ignore', sparse_output=False))
            ]), cat_cols))

        pre = ColumnTransformer(transformers)
        self.pipeline = Pipeline([('pre', pre), ('model', Ridge(alpha=10.0))])

        self.X_train, self.X_test, self.y_train, self.y_test = train_test_split(
            self.X, self.y, test_size=self.test_size, random_state=self.random_state
        )
        self.pipeline.fit(self.X_train, self.y_train)
        self.train_pred = self.pipeline.predict(self.X_train)
        self.test_pred = self.pipeline.predict(self.X_test)

        cv = KFold(n_splits=min(5, len(self.X_train)), shuffle=True, random_state=self.random_state)
        self.cv_r2 = float(cross_val_score(self.pipeline, self.X_train, self.y_train, cv=cv, scoring='r2').mean())

        feature_names = self.pipeline.named_steps['pre'].get_feature_names_out()
        coef = np.abs(np.ravel(self.pipeline.named_steps['model'].coef_))
        self.feature_importance_df = pd.DataFrame({
            'Feature': [f.split('__')[-1] for f in feature_names],
            'Importance': coef
        }).sort_values('Importance', ascending=False).head(20)

    def prepare_sheets_data(self):
        self.summary_df = pd.DataFrame({
            'Item': ['Target Column', 'Total Rows', 'Features Count', 'Train Rows', 'Test Rows', 'Algorithm', 'CV Mean R2'],
            'Value': [str(self.target), len(self.df), self.X.shape[1], len(self.X_train), len(self.X_test), 'Ridge(alpha=10.0)', round(self.cv_r2, 4)]
        })
        self.eda_missing_df = pd.DataFrame({
            'Column': self.df.columns,
            'Missing_Count': self.df.isna().sum().values,
            'Missing_%': (self.df.isna().mean() * 100).round(2).values,
            'Dtype': self.df.dtypes.astype(str).values
        })
        self.eda_stats_df = self.df.describe(include='all').transpose().reset_index().rename(columns={'index': 'Column'})
        self.actual_pred_df = self.X_test.reset_index(drop=True).copy()
        self.actual_pred_df['Actual'] = self.y_test.reset_index(drop=True)
        self.actual_pred_df['Predicted'] = self.test_pred
        self.actual_pred_df['Residual'] = self.actual_pred_df['Actual'] - self.actual_pred_df['Predicted']
        self.actual_pred_df['Abs_Error'] = self.actual_pred_df['Residual'].abs()
        self.actual_pred_df['APE_%'] = np.where(
            self.actual_pred_df['Actual'] != 0,
            (self.actual_pred_df['Abs_Error'] / self.actual_pred_df['Actual'].abs()) * 100,
            np.nan
        )
        rmse_train = float(np.sqrt(mean_squared_error(self.y_train, self.train_pred)))
        rmse_test = float(np.sqrt(mean_squared_error(self.y_test, self.test_pred)))
        self.output_metrics_df = pd.DataFrame([
            ['Train', 'Ridge(alpha=10.0)', mean_absolute_error(self.y_train, self.train_pred), rmse_train, r2_score(self.y_train, self.train_pred)],
            ['Test', 'Ridge(alpha=10.0)', mean_absolute_error(self.y_test, self.test_pred), rmse_test, r2_score(self.y_test, self.test_pred)],
            ['CV Mean', 'Ridge(alpha=10.0)', np.nan, np.nan, self.cv_r2]
        ], columns=['Split', 'Model', 'MAE', 'RMSE', 'R2'])

    def save_all_plots(self):
        sns.set_theme(style='whitegrid')
        self.plot_files = []

        def save(fig, name):
            path = os.path.join(self.plots_dir, name)
            fig.tight_layout()
            fig.savefig(path, dpi=180, bbox_inches='tight')
            plt.close(fig)
            self.plot_files.append(path)

        fig, ax = plt.subplots(figsize=(8, 4.5))
        sns.histplot(self.df[self.target], kde=True, ax=ax, color='#1F4E78')
        ax.set_title(f'Target Distribution: {self.target}')
        save(fig, 'target_distribution.png')

        fig, ax = plt.subplots(figsize=(6.5, 6.5))
        sns.scatterplot(x=self.actual_pred_df['Actual'], y=self.actual_pred_df['Predicted'], ax=ax, color='#1F4E78')
        mn = min(self.actual_pred_df['Actual'].min(), self.actual_pred_df['Predicted'].min())
        mx = max(self.actual_pred_df['Actual'].max(), self.actual_pred_df['Predicted'].max())
        ax.plot([mn, mx], [mn, mx], 'r--', lw=1.5)
        ax.set_title('Actual vs Predicted')
        save(fig, 'actual_vs_predicted.png')

        fig, ax = plt.subplots(figsize=(8, 4.5))
        sns.histplot(self.actual_pred_df['Residual'], kde=True, ax=ax, color='#2CA02C')
        ax.axvline(0, color='red', linestyle='--', lw=1.5)
        ax.set_title('Residual Distribution')
        save(fig, 'residual_distribution.png')

        fig, ax = plt.subplots(figsize=(9, 6))
        sns.barplot(data=self.feature_importance_df.head(15), y='Feature', x='Importance', palette='crest', ax=ax)
        ax.set_title('Top Feature Importances')
        save(fig, 'feature_importance.png')

    def write_excel(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Overview'
        ws['B2'] = 'ML Model Report'
        ws.cell(4, 2, 'Metric Summary')
        for idx, row in enumerate(self.output_metrics_df.itertuples(index=False), 5):
            for col_idx, val in enumerate(row, 2):
                ws.cell(idx, col_idx, None if pd.isna(val) else val)
        wb.save(self.output_excel)
        return self.output_excel

    def run(self):
        self.load_data()
        self.build_model()
        self.prepare_sheets_data()
        self.save_all_plots()
        return self.write_excel()


@app.post("/api/inspect")
async def inspect_dataset(file: Optional[UploadFile] = File(None)):
    target_path = os.path.join(DATA_DIR, "paddydataset.csv")
    if file:
        target_path = os.path.join(DATA_DIR, file.filename)
        with open(target_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

    if not os.path.exists(target_path):
        raise HTTPException(status_code=404, detail="No dataset found.")

    df = pd.read_csv(target_path)
    df.columns = [str(c).strip() for c in df.columns]
    numeric_targets = [
        c for c in df.columns 
        if pd.to_numeric(df[c], errors='coerce').notna().sum() > 10 and df[c].nunique() > 2
    ]
    return {
        "filename": os.path.basename(target_path),
        "total_rows": len(df),
        "columns": list(df.columns),
        "numeric_targets": numeric_targets,
        "preview": df.head(5).to_dict(orient="records")
    }


@app.post("/api/train")
async def train_model(
    filename: str = Form("paddydataset.csv"),
    target_column: str = Form(...),
    test_size: float = Form(0.2),
    random_state: int = Form(42)
):
    try:
        writer = ExcelWriterML(
            data_file=filename,
            target=target_column,
            test_size=test_size,
            random_state=random_state
        )
        writer.run()
        plot_names = [os.path.basename(p) for p in sorted(glob.glob(os.path.join(PLOTS_DIR, "*.png")))]
        return {
            "status": "success",
            "summary": writer.summary_df.to_dict(orient="records"),
            "metrics": writer.output_metrics_df.fillna("").to_dict(orient="records"),
            "plots": [f"http://localhost:8000/plots/{name}" for name in plot_names],
            "excel_ready": True
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/api/download-excel")
async def download_excel():
    excel_path = os.path.join(OUTPUT_DIR, "ml_model_report.xlsx")
    if not os.path.exists(excel_path):
        raise HTTPException(status_code=404, detail="Excel report not generated yet.")
    return FileResponse(
        excel_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="ml_model_report.xlsx"
    )