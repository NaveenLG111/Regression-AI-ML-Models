import os
import sys
import glob
import io
import matplotlib
matplotlib.use('Agg')  # Headless backend prevents GUI/thread crashes in Streamlit
import matplotlib.pyplot as plt
import seaborn as sns
import pandas as pd
import numpy as np
import streamlit as st
from PIL import Image

# Scikit-Learn
from sklearn.model_selection import train_test_split, KFold, cross_val_score
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.linear_model import Ridge

# OpenPyXL
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import ScatterChart, BarChart, Reference, Series


# ==========================================
# 1. CORE ML & EXCEL ENGINE
# ==========================================
class ExcelWriterML:
    def __init__(
        self,
        df_input=None,
        file_path=None,
        output_dir='Output',
        plots_dir='Plots',
        output_excel_name='ml_model_report.xlsx',
        target=None,
        test_size=0.2,
        random_state=42
    ):
        base_dir = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else os.getcwd()

        self.output_dir = os.path.join(base_dir, output_dir)
        self.plots_dir = os.path.join(base_dir, plots_dir)
        self.output_excel = os.path.join(self.output_dir, output_excel_name)

        self.file_path = file_path
        self.df_input = df_input
        self.target = target
        self.test_size = test_size
        self.random_state = random_state

        self.df = None
        self.X = None
        self.y = None
        self.X_train = None
        self.X_test = None
        self.y_train = None
        self.y_test = None
        self.pipeline = None
        self.train_pred = None
        self.test_pred = None
        self.cv_r2 = None

        self.summary_df = None
        self.eda_missing_df = None
        self.eda_stats_df = None
        self.actual_pred_df = None
        self.output_metrics_df = None
        self.feature_importance_df = None
        self.numeric_df = None
        self.plot_files = []

        os.makedirs(self.plots_dir, exist_ok=True)
        os.makedirs(self.output_dir, exist_ok=True)
        self._init_styles()

    def _init_styles(self):
        self.hdr_fill = PatternFill('solid', fgColor='1F4E78')
        self.hdr_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        self.body_font = Font(name='Calibri', size=11)
        self.title_font = Font(name='Calibri', size=15, bold=True, color='1F4E78')
        self.center = Alignment(horizontal='center', vertical='center')
        self.left = Alignment(horizontal='left', vertical='center')
        self.right = Alignment(horizontal='right', vertical='center')
        thin = Side(style='thin', color='D9D9D9')
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def load_data(self):
        if self.df_input is not None:
            self.df = self.df_input.copy()
        elif self.file_path and os.path.exists(self.file_path):
            self.df = pd.read_csv(self.file_path)
        else:
            raise FileNotFoundError(f"Input file not found at: {self.file_path}")

        # Clean column names (strip whitespace)
        self.df.columns = [str(c).strip() for c in self.df.columns]

        if not self.target or self.target not in self.df.columns:
            raise ValueError(f"Target column '{self.target}' not found in dataset columns.")

        # Ensure target is purely numeric
        self.df[self.target] = pd.to_numeric(self.df[self.target], errors='coerce')
        
        # Drop rows where target is NaN
        self.df = self.df.dropna(subset=[self.target]).reset_index(drop=True)
        
        if len(self.df) == 0:
            raise ValueError(f"Target column '{self.target}' does not contain any valid numeric values.")

        self.X = self.df.drop(columns=[self.target])
        self.y = self.df[self.target].astype(float)
        self.numeric_df = self.df.select_dtypes(include=np.number)

    def build_model(self):
        num_cols = [c for c in self.X.columns if pd.api.types.is_numeric_dtype(self.X[c])]
        cat_cols = [c for c in self.X.columns if c not in num_cols]

        transformers = []
        if num_cols:
            transformers.append(('num', Pipeline([
                ('imputer', SimpleImputer(strategy='median')),
                ('scaler', StandardScaler())
            ]), num_cols))
            
        if cat_cols:
            transformers.append(('cat', Pipeline([
                ('imputer', SimpleImputer(strategy='most_frequent')),
                ('onehot', OneHotEncoder(handle_unknown='ignore', sparse_output=False))
            ]), cat_cols))

        pre = ColumnTransformer(transformers)
        model = Ridge(alpha=10.0)
        self.pipeline = Pipeline([('pre', pre), ('model', model)])

        self.X_train, self.X_test, self.y_train, self.y_test = train_test_split(
            self.X, self.y, test_size=self.test_size, random_state=self.random_state
        )

        self.pipeline.fit(self.X_train, self.y_train)
        self.train_pred = self.pipeline.predict(self.X_train)
        self.test_pred = self.pipeline.predict(self.X_test)

        cv = KFold(n_splits=min(5, len(self.X_train)), shuffle=True, random_state=self.random_state)
        self.cv_r2 = cross_val_score(self.pipeline, self.X_train, self.y_train, cv=cv, scoring='r2').mean()

        feature_names = self.pipeline.named_steps['pre'].get_feature_names_out()
        coef = np.abs(np.ravel(self.pipeline.named_steps['model'].coef_))
        self.feature_importance_df = pd.DataFrame({
            'Feature': [f.split('__')[-1] for f in feature_names],
            'Importance': coef
        }).sort_values('Importance', ascending=False).head(20)

    def prepare_sheets_data(self):
        self.summary_df = pd.DataFrame({
            'Item': ['Target Column', 'Total Rows', 'Features Count', 'Train Rows', 'Test Rows', 'Algorithm', 'CV Mean R2 (5-Fold)'],
            'Value': [str(self.target), len(self.df), self.X.shape[1], len(self.X_train), len(self.X_test), 'Ridge(alpha=10.0)', round(float(self.cv_r2), 4)]
        })

        self.eda_missing_df = pd.DataFrame({
            'Column': self.df.columns,
            'Missing_Count': self.df.isna().sum().values,
            'Missing_%': (self.df.isna().mean() * 100).round(2).values,
            'Dtype': self.df.dtypes.astype(str).values
        })

        self.eda_stats_df = self.df.describe(include='all').transpose().reset_index().rename(columns={'index': 'Column'})

        self.actual_pred_df = self.X_test.reset_index(drop=True).copy()
        self.actual_pred_df['Actual'] = self.y_test.reset_index(drop=True)
        self.actual_pred_df['Predicted'] = self.test_pred
        self.actual_pred_df['Residual'] = self.actual_pred_df['Actual'] - self.actual_pred_df['Predicted']
        self.actual_pred_df['Abs_Error'] = self.actual_pred_df['Residual'].abs()
        self.actual_pred_df['APE_%'] = np.where(
            self.actual_pred_df['Actual'] != 0,
            (self.actual_pred_df['Abs_Error'] / self.actual_pred_df['Actual'].abs()) * 100,
            np.nan
        )

        rmse_train = float(np.sqrt(mean_squared_error(self.y_train, self.train_pred)))
        rmse_test = float(np.sqrt(mean_squared_error(self.y_test, self.test_pred)))

        self.output_metrics_df = pd.DataFrame([
            ['Train', 'Ridge(alpha=10.0)', mean_absolute_error(self.y_train, self.train_pred), rmse_train, r2_score(self.y_train, self.train_pred)],
            ['Test', 'Ridge(alpha=10.0)', mean_absolute_error(self.y_test, self.test_pred), rmse_test, r2_score(self.y_test, self.test_pred)],
            ['CV Mean', 'Ridge(alpha=10.0)', np.nan, np.nan, self.cv_r2]
        ], columns=['Split', 'Model', 'MAE', 'RMSE', 'R2'])

    def save_all_plots(self):
        sns.set_theme(style='whitegrid')
        self.plot_files = []

        def save(fig, name):
            path = os.path.join(self.plots_dir, name)
            fig.tight_layout()
            fig.savefig(path, dpi=180, bbox_inches='tight')
            plt.close(fig)
            self.plot_files.append(path)

        # 1. Target Distribution
        fig, ax = plt.subplots(figsize=(8, 4.5))
        sns.histplot(self.df[self.target], kde=True, ax=ax, color='#1F4E78')
        ax.set_title(f'Target Distribution: {self.target}', fontsize=12, fontweight='bold')
        save(fig, 'target_distribution.png')

        # 2. Correlation Heatmap
        if len(self.numeric_df.columns) >= 2:
            fig, ax = plt.subplots(figsize=(9, 7))
            corr = self.numeric_df.corr(numeric_only=True)
            sns.heatmap(corr, cmap='vlag', center=0, annot=False, ax=ax)
            ax.set_title('Correlation Heatmap', fontsize=12, fontweight='bold')
            save(fig, 'correlation_heatmap.png')

        # 3. First Feature vs Target
        num_features = [c for c in self.X.columns if pd.api.types.is_numeric_dtype(self.X[c])]
        if num_features:
            feat = num_features[0]
            fig, ax = plt.subplots(figsize=(8, 4.5))
            sns.scatterplot(x=self.df[feat], y=self.df[self.target], ax=ax, color='#1F4E78', alpha=0.7)
            ax.set_title(f'{feat} vs {self.target}', fontsize=12, fontweight='bold')
            save(fig, 'feature_vs_target.png')

        # 4. Actual vs Predicted
        fig, ax = plt.subplots(figsize=(6.5, 6.5))
        sns.scatterplot(x=self.actual_pred_df['Actual'], y=self.actual_pred_df['Predicted'], ax=ax, color='#1F4E78', alpha=0.7)
        mn = min(self.actual_pred_df['Actual'].min(), self.actual_pred_df['Predicted'].min())
        mx = max(self.actual_pred_df['Actual'].max(), self.actual_pred_df['Predicted'].max())
        ax.plot([mn, mx], [mn, mx], 'r--', lw=1.5, label='Ideal 1:1')
        ax.legend()
        ax.set_title('Actual vs Predicted', fontsize=12, fontweight='bold')
        save(fig, 'actual_vs_predicted.png')

        # 5. Residual Distribution
        fig, ax = plt.subplots(figsize=(8, 4.5))
        sns.histplot(self.actual_pred_df['Residual'], kde=True, ax=ax, color='#2CA02C')
        ax.axvline(0, color='red', linestyle='--', lw=1.5)
        ax.set_title('Residual Distribution', fontsize=12, fontweight='bold')
        save(fig, 'residual_distribution.png')

        # 6. Residuals vs Predicted
        fig, ax = plt.subplots(figsize=(8, 4.5))
        sns.scatterplot(x=self.actual_pred_df['Predicted'], y=self.actual_pred_df['Residual'], ax=ax, alpha=0.7)
        ax.axhline(0, color='red', linestyle='--', lw=1.5)
        ax.set_title('Residuals vs Fitted (Predicted)', fontsize=12, fontweight='bold')
        save(fig, 'residuals_vs_predicted.png')

        # 7. Model R2
        fig, ax = plt.subplots(figsize=(5.5, 4))
        sns.barplot(data=self.output_metrics_df[self.output_metrics_df['Split'] != 'CV Mean'], x='Split', y='R2', palette='Blues_r', ax=ax)
        ax.set_ylim(0, 1.05)
        ax.set_title('Train vs Test R²', fontsize=12, fontweight='bold')
        save(fig, 'train_test_r2.png')

        # 8. Feature Importance
        fig, ax = plt.subplots(figsize=(9, 6))
        sns.barplot(data=self.feature_importance_df.head(15), y='Feature', x='Importance', palette='crest', ax=ax)
        ax.set_title('Top 15 Feature Importances (|Coefficient|)', fontsize=12, fontweight='bold')
        save(fig, 'feature_importance.png')

    def _write_df(self, ws, df, start_row=2, start_col=2, table_name='Table1'):
        for j, c in enumerate(df.columns, start_col):
            cell = ws.cell(start_row, j, c)
            cell.fill = self.hdr_fill
            cell.font = self.hdr_font
            cell.alignment = self.center
            cell.border = self.border

        for i, row in enumerate(df.itertuples(index=False), start_row + 1):
            for j, val in enumerate(row, start_col):
                cell = ws.cell(i, j, None if pd.isna(val) else val)
                cell.font = self.body_font
                cell.border = self.border
                if isinstance(val, (int, np.integer)):
                    cell.alignment = self.right
                    cell.number_format = '#,##0'
                elif isinstance(val, (float, np.floating)):
                    cell.alignment = self.right
                    cell.number_format = '#,##0.0000'
                else:
                    cell.alignment = self.left

        end_row = start_row + len(df)
        end_col = start_col + len(df.columns) - 1

        ref = f'{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}'
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)

        for c in range(start_col, end_col + 1):
            col_letter = get_column_letter(c)
            max_len = 0
            for r in range(start_row, min(end_row, start_row + 30) + 1):
                val = ws.cell(r, c).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(13, min(32, max_len + 3))

        ws.freeze_panes = f'{get_column_letter(start_col)}{start_row + 1}'
        return end_row, end_col

    def _add_actual_pred_chart(self, ws, start_row, actual_col, predicted_col, end_row):
        chart = ScatterChart()
        chart.title = 'Actual vs Predicted (Test Set)'
        chart.style = 13
        chart.x_axis.title = 'Actual'
        chart.y_axis.title = 'Predicted'

        x = Reference(ws, min_col=actual_col, min_row=start_row + 1, max_row=end_row)
        y = Reference(ws, min_col=predicted_col, min_row=start_row + 1, max_row=end_row)
        series = Series(y, x, title='Predictions')
        series.marker.symbol = 'circle'
        series.marker.size = 5
        series.graphicalProperties.line.noFill = True

        chart.series.append(series)
        chart.width = 16
        chart.height = 10
        ws.add_chart(chart, f'{get_column_letter(predicted_col + 5)}2')

    def _add_output_chart(self, ws, start_row, end_row):
        chart = BarChart()
        chart.title = 'Model Performance (R² Score)'
        chart.style = 10
        chart.y_axis.title = 'R²'
        chart.x_axis.title = 'Data Split'

        data = Reference(ws, min_col=5, min_row=start_row, max_row=end_row)
        cats = Reference(ws, min_col=2, min_row=start_row + 1, max_row=end_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        chart.width = 14
        chart.height = 8
        ws.add_chart(chart, 'H2')

    def write_excel(self):
        wb = Workbook()

        ws_overview = wb.active
        ws_overview.title = 'Overview'
        ws_eda = wb.create_sheet('EDA')
        ws_avp = wb.create_sheet('Actuals_vs_Predicted')
        ws_output = wb.create_sheet('Output')

        for ws in wb.worksheets:
            ws.column_dimensions['A'].width = 3

        ws_overview['B2'] = 'ML Model Execution Report'
        ws_overview['B2'].font = self.title_font
        self._write_df(ws_overview, self.summary_df, 4, 2, 'OverviewTbl')

        r, _ = self._write_df(ws_eda, self.eda_missing_df, 2, 2, 'EDAMissingTbl')
        self._write_df(ws_eda, self.eda_stats_df.fillna(''), r + 3, 2, 'EDAStatsTbl')

        ar, _ = self._write_df(ws_avp, self.actual_pred_df, 2, 2, 'ActualPredTbl')
        actual_col = list(self.actual_pred_df.columns).index('Actual') + 2
        predicted_col = list(self.actual_pred_df.columns).index('Predicted') + 2
        self._add_actual_pred_chart(ws_avp, 2, actual_col, predicted_col, ar)

        orow, _ = self._write_df(ws_output, self.output_metrics_df, 2, 2, 'OutputTbl')
        self._add_output_chart(ws_output, 2, orow)

        wb.save(self.output_excel)
        return self.output_excel

    def run(self):
        self.load_data()
        self.build_model()
        self.prepare_sheets_data()
        self.save_all_plots()
        return self.write_excel()


# ==========================================
# 2. STREAMLIT FRONTEND APPLICATION
# ==========================================
def main():
    st.set_page_config(
        page_title="Paddy ML Prediction & Excel Reporter",
        page_icon="🌾",
        layout="wide"
    )

    st.title("🌾 Paddy ML Prediction & Excel Analytics")
    st.caption("End-to-End Regression Pipeline with Exploratory Data Analysis & Excel Generation")

    base_dir = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else os.getcwd()
    data_dir = os.path.join(base_dir, 'Data')
    output_dir = os.path.join(base_dir, 'Output')
    plots_dir = os.path.join(base_dir, 'Plots')

    os.makedirs(data_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(plots_dir, exist_ok=True)

    st.sidebar.header("⚙️ Configuration")
    
    default_csv_path = os.path.join(data_dir, 'paddydataset.csv')
    df = None

    uploaded_file = st.sidebar.file_uploader("Upload CSV Dataset", type=['csv'])

    if uploaded_file is not None:
        try:
            df = pd.read_csv(uploaded_file)
            st.sidebar.success(f"Uploaded: {uploaded_file.name}")
        except Exception as e:
            st.sidebar.error(f"Error reading CSV: {e}")
    elif os.path.exists(default_csv_path):
        try:
            df = pd.read_csv(default_csv_path)
            st.sidebar.info("Loaded default dataset from `Data/paddydataset.csv`")
        except Exception as e:
            st.sidebar.error(f"Error reading default CSV: {e}")

    if df is None:
        st.warning("⚠️ No dataset found. Please upload `paddydataset.csv` using the left sidebar to proceed.")
        st.stop()

    df.columns = [str(c).strip() for c in df.columns]

    # Filter to only numeric columns for regression targets
    numeric_columns = [
        c for c in df.columns 
        if pd.to_numeric(df[c], errors='coerce').notna().sum() > 10 and df[c].nunique() > 2
    ]

    if not numeric_columns:
        st.error("No numeric columns found suitable for regression target.")
        st.stop()

    # Prioritize 'Nursery area (Cents)' or similar numeric target keywords
    default_idx = 0
    for idx, c in enumerate(numeric_columns):
        if 'nursery' in c.lower() or 'cents' in c.lower() or 'yield' in c.lower() or 'production' in c.lower():
            default_idx = idx
            break

    target_column = st.sidebar.selectbox(
        "Select Target Variable (Numeric)",
        options=numeric_columns,
        index=default_idx,
        help="Only numeric columns are available for regression."
    )
    
    test_size = st.sidebar.slider("Test Size Ratio", min_value=0.1, max_value=0.4, value=0.2, step=0.05)
    random_seed = st.sidebar.number_input("Random State Seed", value=42, step=1)

    if st.sidebar.button("🚀 Train Model & Generate Excel", type="primary"):
        with st.spinner("Executing pipeline, creating charts, and generating Excel workbook..."):
            try:
                engine = ExcelWriterML(
                    df_input=df,
                    output_dir='Output',
                    plots_dir='Plots',
                    output_excel_name='ml_model_report.xlsx',
                    target=target_column,
                    test_size=test_size,
                    random_state=int(random_seed)
                )
                excel_file = engine.run()

                st.session_state['engine'] = engine
                st.session_state['excel_file'] = excel_file
                st.success("✅ Model training and Excel report generation complete!")
            except Exception as e:
                st.error(f"Execution Error: {e}")
                st.stop()

    if 'engine' in st.session_state:
        engine = st.session_state['engine']
        excel_file = st.session_state['excel_file']

        tab1, tab2, tab3, tab4 = st.tabs(["📊 Performance & Summary", "📈 Visualization Plots", "📋 Predictions Sample", "💾 Export Excel"])

        with tab1:
            col1, col2 = st.columns(2)
            with col1:
                st.subheader("Model Summary")
                st.dataframe(engine.summary_df, use_container_width=True, hide_index=True)
            with col2:
                st.subheader("Performance Metrics")
                st.dataframe(engine.output_metrics_df, use_container_width=True, hide_index=True)

        with tab2:
            st.subheader("Generated Visual Plots")
            plot_files = sorted(glob.glob(os.path.join(engine.plots_dir, "*.png")))
            if plot_files:
                for i in range(0, len(plot_files), 2):
                    c1, c2 = st.columns(2)
                    c1.image(plot_files[i], use_container_width=True)
                    if i + 1 < len(plot_files):
                        c2.image(plot_files[i + 1], use_container_width=True)
            else:
                st.info("No plots generated.")

        with tab3:
            st.subheader("Test Set: Actual vs Predicted Sample")
            preview_cols = ['Actual', 'Predicted', 'Residual', 'APE_%']
            feature_cols = [c for c in engine.actual_pred_df.columns if c not in preview_cols][:4]
            st.dataframe(engine.actual_pred_df[feature_cols + preview_cols].head(25), use_container_width=True)

        with tab4:
            st.subheader("Download Generated Excel Report")
            if os.path.exists(excel_file):
                with open(excel_file, "rb") as f:
                    file_bytes = f.read()
                st.download_button(
                    label="📥 Download ml_model_report.xlsx",
                    data=file_bytes,
                    file_name=os.path.basename(excel_file),
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
    else:
        st.info("👈 Review your dataset settings in the left sidebar and click **'Train Model & Generate Excel'**.")
        st.subheader("Dataset Preview")
        st.dataframe(df.head(10), use_container_width=True)


if __name__ == '__main__':
    main()